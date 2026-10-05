import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    navy = RGBColor(0x1E, 0x3A, 0x5F)
    sky_blue = RGBColor(0x02, 0x84, 0xC7)
    dark_gray = RGBColor(0x33, 0x41, 0x55)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("MINI-PROJECT SHORT TECHNICAL REPORT")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    # Subtitle / Metadata
    meta_lines = [
        ("Course: ", "Cross-Platform Mobile App Development (VKU)"),
        ("Mini-Project Title: ", "Mini-Project 2: Real-time Study Room Booking App (React Native & Expo)"),
        ("Team / Student Name: ", "Lê Cảm (CAMLC25)"),
        ("Submission Date: ", "05/10/2026"),
    ]
    for label, val in meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run(label)
        r_lbl.bold = True
        r_lbl.font.color.rgb = navy
        r_val = p.add_run(val)
        r_val.font.color.rgb = dark_gray

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("1. GENERAL INFORMATION & DELIVERABLE LINKS")
    r.font.color.rgb = navy

    items_1 = [
        ("Team Members: ", "Lê Cảm — Student ID: 21IT001 — Role: Fullstack Mobile Architecture & Concurrency Engine — Contribution: 100%"),
        ("🔗 Live Demo URL: ", "https://camle-vku-study-room.lecam.workers.dev"),
        ("💻 GitHub Repository: ", "https://github.com/CAMLC25/camle-vku-study-room"),
        ("🎥 Video Demo (Optional): ", "https://camle-vku-study-room.lecam.workers.dev"),
    ]
    for label, text in items_1:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(label)
        r1.bold = True
        r1.font.color.rgb = sky_blue if "URL" in label or "Repo" in label else navy
        r2 = p.add_run(text)
        r2.font.color.rgb = dark_gray

    # Section 2
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run("2. FEATURE IMPLEMENTATION CHECKLIST")
    r.font.color.rgb = navy

    features = [
        ("1", "Responsive Viewport & Safe Area", "✅ Complete", "Custom hook useResponsiveLayout adapts across Mobile (1 col) and Desktop/Tablet (2 cols) with Math.min(width, 1140) width clamping and flex: 1 columns preventing offscreen clipping. Full dynamic inset handling via react-native-safe-area-context."),
        ("2", "Server State & Caching (TanStack Query)", "✅ Complete", "Integrated @tanstack/react-query with QueryClientProvider (staleTime: 5min, gcTime: 10min). Custom hook useRooms() with pull-to-refresh on FlatList."),
        ("3", "Client State & Persistence (Zustand)", "✅ Complete", "Zustand store with persist middleware storing in @react-native-async-storage/async-storage under key 'vku-booking-storage'."),
        ("4", "Type-Safe Navigation (React Navigation 7)", "✅ Complete", "Nested Stack + Bottom Tabs architecture (RootNavigator + MainTabs) with strict TypeScript types (RootStackParamList, MainTabParamList)."),
        ("5", "Instant Search & Multi-Parameter Filter", "✅ Complete", "Real-time text search and filter chips by Building (A, B, C, V), Capacity (2–20), and Equipment (Projector, Whiteboard, High-spec PC, AC)."),
        ("6", "7-Day Rolling Horizon & Discrete Slots", "✅ Complete", "Interactive date selector for 7 days ahead; 4 discrete 2-hour daily sessions (07:30–09:30, 09:30–11:30, 13:00–15:00, 15:00–17:00)."),
        ("7", "90-Second Soft Hold & Realtime Sync", "✅ Complete", "Instant reservation lock with 90s countdown modal; broadcasts HELD_BY_OTHER amber state to other peers via Supabase Realtime channels."),
        ("8", "Concurrency & Anti-Collision Engine", "✅ Complete", "Server-authoritative PostgreSQL advisory locking and partial unique index (idx_bookings_active_slot). Client-generated UUID tokens prevent replay duplicates."),
        ("9", "Offline Outbox & Conflict Resolution", "✅ Complete", "Offline booking requests queued in an outbox (PENDING_SYNC) and flushed sequentially upon reconnection with deterministic conflict handling."),
        ("10", "QR Check-in Pass & Local Notifications", "✅ Complete", "Interactive check-in pass powered by react-native-qrcode-svg; scheduled 15-minute countdown reminders via expo-notifications."),
        ("11", "Layout Animations (Reanimated 3/4)", "✅ Complete", "Staggered entrance animations on room feed using FadeInDown.delay(index * 60).springify() and spring physics on interactions."),
        ("12", "Student Auth, Email Verification & Profile Sync", "✅ Complete", "Full authentication with real email/password via Supabase Auth (auth.users), custom Gmail SMTP Relay bypassing free-tier rate limits, mandatory email confirmation with 1-tap deep link auto-login (#access_token), and PostgreSQL trigger on auth.users syncing to public.students."),
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["#", "Required Feature", "Status", "Implementation Details & Acceptance Level"]
    col_widths = [Inches(0.4), Inches(2.2), Inches(1.1), Inches(3.1)]

    for idx, (cell, title, w) in enumerate(zip(hdr_cells, hdr_titles, col_widths)):
        cell.width = w
        set_cell_background(cell, "1E3A5F")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    for row_idx, (num, feat, stat, desc) in enumerate(features):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for idx, (cell, val, w) in enumerate(zip(row_cells, [num, feat, stat, desc], col_widths)):
            cell.width = w
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.size = Pt(9)
            if idx == 1:
                run.bold = True
                run.font.color.rgb = navy
            elif idx == 2:
                run.bold = True
                run.font.color.rgb = RGBColor(0x10, 0xB9, 0x81)
            else:
                run.font.color.rgb = dark_gray

    # Section 3
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run("3. TECHNICAL ARCHITECTURE & PROJECT STRUCTURE")
    r.font.color.rgb = navy

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("3.1 Directory Organization")
    r.bold = True
    r.font.color.rgb = sky_blue

    tree_text = (
        "vku-study-room/\n"
        "├── App.tsx                     # Entrypoint with Safe Area, QueryProvider & RootNavigator\n"
        "├── app.json                    # Expo Managed configuration & asset definitions\n"
        "├── wrangler.jsonc              # Cloudflare Workers Static Assets deployment config\n"
        "├── src/\n"
        "│   ├── components/             # Reusable UI components (RoomCard, SearchBar, FilterChips, Modal)\n"
        "│   ├── config/                 # Environment variables and runtime mode detection\n"
        "│   ├── hooks/                  # Custom hooks (useRooms, useResponsiveLayout, useRoomAvailability)\n"
        "│   ├── i18n/                   # Multi-language dictionary (Vietnamese & English)\n"
        "│   ├── navigation/             # Type-safe navigators (RootNavigator, MainTabs, types.ts)\n"
        "│   ├── providers/              # TanStack QueryClientProvider configuration\n"
        "│   ├── screens/                # Core screens (RoomListScreen, RoomDetailScreen, MyBookingsScreen, ProfileScreen, LoginScreen, RegisterScreen)\n"
        "│   ├── services/               # Abstraction layer (Supabase, Mock, Outbox Sync, Notifications, Auth)\n"
        "│   ├── store/                  # Zustand stores (useBookingStore with persist, useAuthStore, useLanguageStore)\n"
        "│   ├── theme/                  # Design tokens, typography, and VKU color palettes\n"
        "│   └── types/                  # Strict TypeScript domain interfaces\n"
        "├── supabase/\n"
        "│   ├── schema.sql              # PostgreSQL DDL, advisory lock RPC, and RLS policies\n"
        "│   └── seed.sql                # 20 pre-seeded study rooms across 4 VKU campus buildings\n"
        "└── test/                       # Node.js automated verification harness (Concurrency, Outbox, Auth, QR)"
    )
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_after = Pt(6)
    r_code = p_code.add_run(tree_text)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8)
    r_code.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("3.2 State Management & Data Flow Architecture")
    r.bold = True
    r.font.color.rgb = sky_blue

    state_desc = [
        ("Server State (TanStack Query): ", "Manages remote entity caches ('rooms') with a 5-minute staleTime, automated background synchronization, and zero-redundancy network refetching on pull-to-refresh."),
        ("Client State (Zustand): ", "Manages volatile and local domain state (active reservations, active 90s hold session, outbox queue, filter parameters, and user authentication). All critical mutations persist immediately to device storage via @react-native-async-storage/async-storage ('vku-booking-storage')."),
    ]
    for lbl, desc in state_desc:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(lbl)
        r1.bold = True
        r1.font.color.rgb = navy
        r2 = p.add_run(desc)
        r2.font.color.rgb = dark_gray

    # Section 4
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run("4. EMPIRICAL EVIDENCE & SCREENSHOTS")
    r.font.color.rgb = navy

    screens = [
        ("Room Discovery & Multi-Parameter Filter (RoomListScreen): ", "Displays 20 study rooms across Buildings A, B, C, and V in an adaptive 60 FPS grid. Features real-time filtering chips by building, seat capacity (2–20), and equipment tags with smooth staggered entry animations (FadeInDown.springify()). Dynamic column calculation with 1140px max-width clamping ensures zero cut-off cards on desktop viewports."),
        ("Interactive Time-Slot Selector & 90s Countdown Modal (RoomDetailScreen): ", "Features a 7-day rolling calendar with 4 discrete two-hour slots per day. Selecting an available slot initiates a 90-second soft hold with a synchronized countdown timer, broadcasting HELD_BY_OTHER amber badges across all connected devices."),
        ("Personal Reservations & Offline Synchronization (MyBookingsScreen): ", "Segmented tabs for 'All', 'Confirmed', 'Pending Sync', 'Conflicted', and 'Cancelled'. Displays local outbox queue status during offline mode and automatically resolves conflicts with optimistic feedback."),
        ("Digital QR Booking Pass (Check-in Modal): ", "Generates a high-contrast, cryptographically clean QR code encoding the unique booking token for campus security check-in."),
        ("Modern Minimalist Authentication (LoginScreen & RegisterScreen): ", "Streamlined, distraction-free authentication interface enforcing valid email credentials. Seamless handling of email confirmation tokens via Supabase Auth deep links with automated URL hash sanitization."),
    ]
    for lbl, desc in screens:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(lbl)
        r1.bold = True
        r1.font.color.rgb = navy
        r2 = p.add_run(desc)
        r2.font.color.rgb = dark_gray

    # Section 5
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)
    r = h5.add_run("5. TECHNICAL CHALLENGES & RESOLUTIONS")
    r.font.color.rgb = navy

    challenges = [
        ("5.1 Challenge 1 (State Management & Caching): Strict Decoupling of Client State (Zustand) & Server State (TanStack Query)",
         "Problem: In accordance with the VKU Week 6 curriculum guideline ('Zustand for client state + TanStack Query for server state — don't mix them'), a major architecture challenge was maintaining a clean separation of concerns. Storing rooms and remote availability entirely in Zustand led to stale data, memory overhead, and lacked automated background revalidation and pull-to-refresh. Conversely, managing volatile checkout holds (90s countdown), search filters, and persistent offline outbox mutations in TanStack Query degraded UI immediacy and prevented robust @react-native-async-storage/async-storage persistence under key 'vku-booking-storage'.\n"
         "Resolution: (1) Server State (TanStack Query 5): Encapsulated the remote room catalog inside a custom useRooms(building) hook, configured with staleTime: 5min and gcTime: 10min. Wired data, isLoading, and refetch directly into FlatList's native refreshing and onRefresh props for pull-to-refresh capabilities. (2) Client State (Zustand 5 + persist): Confined user reservations (myBookings), offline outbox, active 90s hold sessions, and UI filter chips to useBookingStore. Enforced AsyncStorage persistence under key 'vku-booking-storage' with narrow, atomic selector subscriptions (useBookingStore(s => s.myBookings)), eliminating unnecessary full-screen re-renders."),
        ("5.2 Challenge 2 (Core Features & Concurrency): Eliminating Race Conditions (TOCTOU) & Managing 90-Second Soft Holds",
         "Problem: In high-concurrency academic settings (e.g. final exam revision week), multiple students attempt to book the exact same study room and time slot simultaneously. Naive client-side availability checks (SELECT ... WHERE status = 'CONFIRMED') suffer from Time-of-Check to Time-of-Use (TOCTOU) race conditions, resulting in catastrophic double-bookings. Furthermore, if a student opens a hold modal and cancels or disconnects mid-flight, unreleased holds create lingering 'phantom locks'.\n"
         "Resolution: (1) Realtime 90-Second Soft Hold: Instantiates an ephemeral reservation hold with a synchronized countdown timer upon slot selection, broadcasting HELD_BY_OTHER amber badges to all connected peers via Supabase Realtime channels. An isCancelledRef guard ensures that if a user cancels the modal before the server responds, an immediate background releaseHold is dispatched to return the slot to AVAILABLE. (2) PostgreSQL Advisory Locking & Partial Unique Index: At the persistence layer, the stored procedure book_slot() enforces exclusive transaction-level advisory locks via PERFORM pg_advisory_xact_lock(hashtext(p_room_id || ':' || p_booking_date || ':' || p_slot_index)), backed by a partial unique index idx_bookings_active_slot (room_id, booking_date, slot_index) WHERE status = 'CONFIRMED'. Any concurrent attempt hitting the exact same microsecond fails with PostgreSQL error 23505 (SLOT_ALREADY_BOOKED). (3) Client-side Idempotency Keys (UUID v4): Protects against duplicate bookings caused by mobile packet loss and automatic HTTP retry flings."),
        ("5.3 Challenge 3 (UI/UX & Performance): 60 FPS FlatList Virtualization, Reanimated Layout Animations & Desktop Grid Clamping",
         "Problem: Rendering a catalog of 20+ study rooms containing remote photographic thumbnails, status tags, and equipment badges causes frame drops and layout stutter on mobile devices during momentum flinging. Additionally, when running across cross-platform viewports (Web desktop 1920px+ vs Mobile), naive card width calculation using raw window.width pushed the second column of cards offscreen beyond the centered maxWidth: 1140px container.\n"
         "Resolution: (1) FlatList Virtualization & Sync Layout: Fixed item dimensions to ROOM_CARD_HEIGHT = 136px and provided getItemLayout for synchronous coordinate calculation without bridge measurement. Configured initialNumToRender: 8, maxToRenderPerBatch: 8, windowSize: 5, and removeClippedSubviews: true. (2) Worklet-Driven Layout Animations (Reanimated 3): Integrated FadeInDown.delay(index * 60).springify() executing directly on the UI thread via Hermes worklets, completely bypassing the asynchronous JavaScript bridge for buttery 60/120 FPS transitions. (3) Adaptive Width Clamping: Re-architected useResponsiveLayout to clamp calculation to effectiveWidth = Math.min(width, 1140). In RoomListScreen, each card container is styled with flex: 1 and maxWidth: cardWidth, guaranteeing symmetric 2-column distribution and 0% card cut-off across all display form factors."),
        ("5.4 Challenge 4 (Navigation Architecture & Type Safety): Nested Type-Safe Routing (React Navigation 7) & Deep Linking URL Sanitization",
         "Problem: Mini-Project 2 requires a complex nested navigation hierarchy combining Root Native Stack, Bottom Tabs (BrowseRooms, MyBookings, Profile), and Modal presentations (ConfirmBookingModal, BookingPassModal). In loosely-typed setups, route typos and missing route parameters produce silent runtime crashes on mobile devices. Additionally, email verification deep links deposit access tokens into the web address bar (#access_token=...), triggering unintended authentication reload loops on refresh.\n"
         "Resolution: (1) Comprehensive TypeScript Route Contracts: Defined strict RootStackParamList and MainTabParamList types, typed every screen with NativeStackScreenProps, and provided typed navigation hooks (useNavigation<NavigationProp>()), guaranteeing 100% compile-time route verification (npx tsc --noEmit). (2) Nesting Architecture: Root Stack wraps Main Tabs, allowing Room Details and Modals to push on top and cleanly hide the tab bar when focused interaction is required. (3) URL Hash Sanitization: Integrated window.history.replaceState into useAuthStore to automatically strip hash fragments (#access_token=...) immediately upon session ingestion and logout, preventing reload loops."),
        ("5.5 Challenge 5 (Offline Outbox & Network Resilience): Sequential Outbox Flusher & Deterministic Conflict Reconciliation",
         "Problem: University campus environments present frequent connectivity dead zones (elevators, basements, crowded lecture halls). Naive mobile implementations either freeze the UI with blocking spinners or mistakenly mark offline bookings as CONFIRMED locally, causing severe disputes when a student arrives at a room that was legitimately booked by someone else online.\n"
         "Resolution: (1) Zero False Confirmation Contract: Bookings initiated while offline are strictly stored in myBookings and outbox with status PENDING_SYNC. The UI displays an amber badge and disables the QR Check-in Pass until server verification is achieved. (2) Sequential Flusher (FIFO): syncService processes pending outbox requests one-by-one via for...of loops with transactional delays (80ms), strictly avoiding Promise.all() to preserve causality and protect student daily quota constraints (max 2 slots/day). (3) Deterministic Conflict Handling: If a slot was taken online while the device was disconnected, the server rejects the synchronization with SLOT_ALREADY_BOOKED. The flusher marks both the outbox item and local booking as CONFLICTED, captures the error message, and dispatches a local high-priority notification via expo-notifications alerting the student to select an alternate slot.")
    ]

    for title, desc in challenges:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.bold = True
        r1.font.color.rgb = sky_blue

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        r2 = p2.add_run(desc)
        r2.font.color.rgb = dark_gray

    output_path = "docs/MINI_PROJECT_2_REPORT.docx"
    doc.save(output_path)
    print("SUCCESS: Generated", output_path)

if __name__ == "__main__":
    create_report()
